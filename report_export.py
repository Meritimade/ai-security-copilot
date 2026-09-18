from __future__ import annotations

import io
import json
from pathlib import Path
from typing import Any

import pandas as pd


def _normalise_list(value):
    if value is None:
        return []

    if isinstance(value, list):
        return [str(item) for item in value]

    return [str(value)]


def build_filtered_package(
    dataset_name: str,
    dataset_type: str,
    dataset_mode: str,
    total_events: int,
    findings: list[dict[str, Any]],
    alerts: list[dict[str, Any]],
    mitre_assessments: list[dict[str, Any]],
    document_result: dict[str, Any] | None = None,
) -> dict[str, Any]:
    package = {
        "report_metadata": {
            "application": "AI Security Copilot",
            "dataset": dataset_name,
            "dataset_type": dataset_type,
            "dataset_mode": dataset_mode,
            "events_analysed": total_events,
        },
        "findings": findings,
        "alerts": alerts,
        "mitre_assessments": mitre_assessments,
    }

    if document_result:
        package["document_analysis"] = document_result

    return package


def package_to_json(package: dict[str, Any]) -> bytes:
    return json.dumps(
        package,
        indent=2,
        ensure_ascii=False,
        default=str,
    ).encode("utf-8")


def package_to_markdown(package: dict[str, Any]) -> bytes:
    meta = package.get("report_metadata", {})
    findings = package.get("findings", [])
    alerts = package.get("alerts", [])
    mitre = package.get("mitre_assessments", [])
    document = package.get("document_analysis")

    lines = [
        "# AI Security Copilot Report",
        "",
        "## Report Scope",
        "",
        f"- Dataset: {meta.get('dataset', '-')}",
        f"- Dataset type: {meta.get('dataset_type', '-')}",
        f"- Dataset mode: {meta.get('dataset_mode', '-')}",
        f"- Events analysed: {meta.get('events_analysed', 0):,}",
        "",
        "## Findings",
        "",
    ]

    if findings:
        for finding in findings:
            name = finding.get("label", finding.get("finding", "Unknown"))
            classification = finding.get("classification", "unknown")
            records = finding.get("record_count", "-")
            share = finding.get("percentage_of_dataset", "-")

            lines.extend([
                f"### {name}",
                f"- Classification: {classification}",
                f"- Records: {records}",
                f"- Dataset share: {share}",
            ])

            limitations = finding.get("evidence_limitations", [])
            if limitations:
                lines.append("- Evidence limitations:")
                lines.extend(f"  - {item}" for item in limitations)

            lines.append("")
    else:
        lines.extend(["No filtered findings.", ""])

    lines.extend([
        "## Alerts",
        "",
    ])

    if alerts:
        for alert in alerts:
            lines.extend([
                f"### {alert.get('alert_id', '-')}: {alert.get('finding', '-')}",
                f"- Severity: {alert.get('severity', '-')}",
                f"- Risk score: {alert.get('risk_score', 0)}",
                f"- Confidence: {alert.get('confidence', '-')}",
                f"- Status: {alert.get('status', '-')}",
                "",
            ])
    else:
        lines.extend(["No filtered alerts.", ""])

    lines.extend([
        "## MITRE ATT&CK Assessment",
        "",
    ])

    if mitre:
        for item in mitre:
            lines.append(
                f"- {item.get('technique_id', '-')}: "
                f"{item.get('technique_name', '-')} — "
                f"{item.get('status', '-')}"
            )
    else:
        lines.append("No filtered MITRE assessments.")

    if document:
        lines.extend([
            "",
            "## Document / Image Analysis",
            "",
            str(document.get("executive_summary", "")),
            "",
        ])

    return "\n".join(lines).encode("utf-8")


def package_to_txt(package: dict[str, Any]) -> bytes:
    return package_to_markdown(package)


def package_to_csv(package: dict[str, Any]) -> bytes:
    rows = []

    for alert in package.get("alerts", []):
        rows.append({
            "Record Type": "Alert",
            "ID": alert.get("alert_id", ""),
            "Finding": alert.get("finding", ""),
            "Severity": alert.get("severity", ""),
            "Risk Score": alert.get("risk_score", ""),
            "Confidence": alert.get("confidence", ""),
            "Status": alert.get("status", ""),
        })

    for item in package.get("mitre_assessments", []):
        rows.append({
            "Record Type": "MITRE",
            "ID": item.get("technique_id", ""),
            "Finding": item.get("finding", ""),
            "Severity": "",
            "Risk Score": "",
            "Confidence": "",
            "Status": item.get("status", ""),
        })

    if not rows:
        rows.append({
            "Record Type": "Summary",
            "ID": "",
            "Finding": "",
            "Severity": "",
            "Risk Score": "",
            "Confidence": "",
            "Status": "No filtered records",
        })

    return pd.DataFrame(rows).to_csv(index=False).encode("utf-8")


def package_to_xlsx(package: dict[str, Any]) -> bytes:
    output = io.BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl",
    ) as writer:

        meta = package.get("report_metadata", {})
        pd.DataFrame(
            [
                {"Metric": "Dataset", "Value": meta.get("dataset", "-")},
                {"Metric": "Dataset Type", "Value": meta.get("dataset_type", "-")},
                {"Metric": "Dataset Mode", "Value": meta.get("dataset_mode", "-")},
                {"Metric": "Events Analysed", "Value": meta.get("events_analysed", 0)},
            ]
        ).to_excel(
            writer,
            sheet_name="Summary",
            index=False,
        )

        alerts = package.get("alerts", [])
        if alerts:
            pd.DataFrame(alerts).to_excel(
                writer,
                sheet_name="Alerts",
                index=False,
            )
        else:
            pd.DataFrame(
                [{"Status": "No filtered alerts"}]
            ).to_excel(
                writer,
                sheet_name="Alerts",
                index=False,
            )

        findings = package.get("findings", [])
        if findings:
            simple_findings = []
            for finding in findings:
                simple_findings.append({
                    "Finding": finding.get(
                        "label",
                        finding.get("finding", "Unknown")
                    ),
                    "Classification": finding.get("classification", ""),
                    "Records": finding.get("record_count", ""),
                    "Dataset Share": finding.get("percentage_of_dataset", ""),
                })

            pd.DataFrame(simple_findings).to_excel(
                writer,
                sheet_name="Findings",
                index=False,
            )
        else:
            pd.DataFrame(
                [{"Status": "No filtered findings"}]
            ).to_excel(
                writer,
                sheet_name="Findings",
                index=False,
            )

        mitre = package.get("mitre_assessments", [])
        if mitre:
            pd.DataFrame(mitre).to_excel(
                writer,
                sheet_name="MITRE",
                index=False,
            )
        else:
            pd.DataFrame(
                [{"Status": "No filtered MITRE assessments"}]
            ).to_excel(
                writer,
                sheet_name="MITRE",
                index=False,
            )

        document = package.get("document_analysis")
        if document:
            pd.DataFrame(
                [{
                    "Executive Summary": document.get(
                        "executive_summary",
                        ""
                    ),
                    "Overall Risk": document.get(
                        "overall_risk",
                        {}
                    ).get("level", "Unknown"),
                    "Confidence": document.get(
                        "overall_risk",
                        {}
                    ).get("confidence", "Low"),
                }]
            ).to_excel(
                writer,
                sheet_name="Document Analysis",
                index=False,
            )

    return output.getvalue()


def package_to_pdf(package: dict[str, Any]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        Table,
        TableStyle,
        PageBreak,
    )

    output = io.BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title="AI Security Copilot Report",
    )

    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="SmallMuted",
            parent=styles["BodyText"],
            fontSize=8,
            leading=11,
            textColor=colors.HexColor("#667085"),
        )
    )

    story = [
        Paragraph(
            "AI Security Copilot Report",
            styles["Title"],
        ),
        Spacer(1, 8),
    ]

    meta = package.get("report_metadata", {})

    scope_table = Table([
        ["Dataset", str(meta.get("dataset", "-"))],
        ["Dataset Type", str(meta.get("dataset_type", "-"))],
        ["Dataset Mode", str(meta.get("dataset_mode", "-"))],
        ["Events Analysed", f"{int(meta.get('events_analysed', 0)):,}"],
    ], colWidths=[45 * mm, 125 * mm])

    scope_table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#eef2f7")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
            ("FONTNAME", (1, 0), (1, -1), "Helvetica"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ])
    )

    story.extend([
        Paragraph("Report Scope", styles["Heading2"]),
        scope_table,
        Spacer(1, 10),
        Paragraph("Findings", styles["Heading2"]),
    ])

    findings = package.get("findings", [])

    if findings:
        finding_rows = [
            ["Finding", "Classification", "Records", "Share"]
        ]

        for finding in findings:
            finding_rows.append([
                str(finding.get(
                    "label",
                    finding.get("finding", "Unknown")
                )),
                str(finding.get("classification", "")),
                str(finding.get("record_count", "")),
                str(finding.get("percentage_of_dataset", "")),
            ])

        table = Table(
            finding_rows,
            repeatRows=1,
            colWidths=[65 * mm, 40 * mm, 30 * mm, 35 * mm],
        )

        table.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#cbd5e1")),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1),
                 [colors.white, colors.HexColor("#f8fafc")]),
            ])
        )

        story.extend([table, Spacer(1, 10)])
    else:
        story.extend([
            Paragraph("No filtered findings.", styles["BodyText"]),
            Spacer(1, 10),
        ])

    story.append(Paragraph("Alerts", styles["Heading2"]))

    alerts = package.get("alerts", [])

    if alerts:
        alert_rows = [
            ["Finding", "Severity", "Risk", "Confidence", "Status"]
        ]

        for alert in alerts:
            alert_rows.append([
                str(alert.get("finding", "")),
                str(alert.get("severity", "")),
                str(alert.get("risk_score", "")),
                str(alert.get("confidence", "")),
                str(alert.get("status", "")),
            ])

        table = Table(
            alert_rows,
            repeatRows=1,
            colWidths=[60 * mm, 30 * mm, 20 * mm, 30 * mm, 30 * mm],
        )

        table.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#cbd5e1")),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1),
                 [colors.white, colors.HexColor("#f8fafc")]),
            ])
        )

        story.extend([table, Spacer(1, 10)])
    else:
        story.extend([
            Paragraph("No filtered alerts.", styles["BodyText"]),
            Spacer(1, 10),
        ])

    story.append(Paragraph("MITRE ATT&CK Assessment", styles["Heading2"]))

    mitre = package.get("mitre_assessments", [])

    if mitre:
        mitre_rows = [["Finding", "Technique", "Name", "Status"]]

        for item in mitre:
            mitre_rows.append([
                str(item.get("finding", "")),
                str(item.get("technique_id", "")),
                str(item.get("technique_name", "")),
                str(item.get("status", "")),
            ])

        table = Table(
            mitre_rows,
            repeatRows=1,
            colWidths=[45 * mm, 25 * mm, 65 * mm, 35 * mm],
        )

        table.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#cbd5e1")),
                ("FONTSIZE", (0, 0), (-1, -1), 6.5),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1),
                 [colors.white, colors.HexColor("#f8fafc")]),
            ])
        )

        story.extend([table, Spacer(1, 10)])
    else:
        story.extend([
            Paragraph("No filtered MITRE assessments.", styles["BodyText"]),
            Spacer(1, 10),
        ])

    document_analysis = package.get("document_analysis")

    if document_analysis:
        story.append(PageBreak())
        story.append(
            Paragraph(
                "Document / Image Analysis",
                styles["Heading2"],
            )
        )

        summary = str(
            document_analysis.get(
                "executive_summary",
                ""
            )
        )

        for paragraph in summary.split("\n"):
            if paragraph.strip():
                story.append(
                    Paragraph(
                        paragraph.replace("&", "&amp;"),
                        styles["BodyText"],
                    )
                )
                story.append(Spacer(1, 4))

        overall = document_analysis.get(
            "overall_risk",
            {}
        )

        story.extend([
            Spacer(1, 8),
            Paragraph(
                f"Overall risk: {overall.get('level', 'Unknown')}",
                styles["Heading3"],
            ),
            Paragraph(
                f"Confidence: {overall.get('confidence', 'Low')}",
                styles["BodyText"],
            ),
            Paragraph(
                str(overall.get("reason", "")),
                styles["BodyText"],
            ),
        ])

    document.build(story)

    return output.getvalue()


EXPORTERS = {
    "PDF": package_to_pdf,
    "DOCX": None,  # implemented below
    "XLSX": package_to_xlsx,
    "CSV": package_to_csv,
    "JSON": package_to_json,
    "TXT": package_to_txt,
    "Markdown": package_to_markdown,
}


def package_to_docx(package: dict[str, Any]) -> bytes:
    from docx import Document

    document = Document()

    document.add_heading("AI Security Copilot Report", 0)

    meta = package.get("report_metadata", {})

    document.add_heading("Report Scope", level=1)

    for label, key in [
        ("Dataset", "dataset"),
        ("Dataset Type", "dataset_type"),
        ("Dataset Mode", "dataset_mode"),
        ("Events Analysed", "events_analysed"),
    ]:
        value = meta.get(key, "-")
        if key == "events_analysed":
            try:
                value = f"{int(value):,}"
            except Exception:
                pass
        document.add_paragraph(f"{label}: {value}")

    document.add_heading("Findings", level=1)

    findings = package.get("findings", [])

    if not findings:
        document.add_paragraph("No filtered findings.")
    else:
        for finding in findings:
            name = finding.get(
                "label",
                finding.get("finding", "Unknown")
            )
            document.add_heading(str(name), level=2)
            document.add_paragraph(
                f"Classification: {finding.get('classification', '')}"
            )
            document.add_paragraph(
                f"Records: {finding.get('record_count', '')}"
            )
            document.add_paragraph(
                f"Dataset share: {finding.get('percentage_of_dataset', '')}"
            )

            limitations = _normalise_list(
                finding.get("evidence_limitations")
            )

            if limitations:
                document.add_paragraph("Evidence limitations:")
                for item in limitations:
                    document.add_paragraph(
                        item,
                        style="List Bullet",
                    )

    document.add_heading("Alerts", level=1)

    alerts = package.get("alerts", [])

    if not alerts:
        document.add_paragraph("No filtered alerts.")
    else:
        table = document.add_table(
            rows=1,
            cols=5,
        )
        headers = [
            "Finding",
            "Severity",
            "Risk Score",
            "Confidence",
            "Status",
        ]

        for cell, header in zip(table.rows[0].cells, headers):
            cell.text = header

        for alert in alerts:
            cells = table.add_row().cells
            values = [
                alert.get("finding", ""),
                alert.get("severity", ""),
                alert.get("risk_score", ""),
                alert.get("confidence", ""),
                alert.get("status", ""),
            ]

            for cell, value in zip(cells, values):
                cell.text = str(value)

    document.add_heading("MITRE ATT&CK Assessment", level=1)

    mitre = package.get("mitre_assessments", [])

    if not mitre:
        document.add_paragraph("No filtered MITRE assessments.")
    else:
        table = document.add_table(
            rows=1,
            cols=4,
        )

        headers = [
            "Finding",
            "Technique",
            "Technique Name",
            "Status",
        ]

        for cell, header in zip(table.rows[0].cells, headers):
            cell.text = header

        for item in mitre:
            cells = table.add_row().cells
            values = [
                item.get("finding", ""),
                item.get("technique_id", ""),
                item.get("technique_name", ""),
                item.get("status", ""),
            ]

            for cell, value in zip(cells, values):
                cell.text = str(value)

    document_analysis = package.get("document_analysis")

    if document_analysis:
        document.add_heading(
            "Document / Image Analysis",
            level=1,
        )

        document.add_paragraph(
            str(
                document_analysis.get(
                    "executive_summary",
                    ""
                )
            )
        )

        overall = document_analysis.get(
            "overall_risk",
            {}
        )

        document.add_paragraph(
            f"Overall risk: {overall.get('level', 'Unknown')}"
        )

        document.add_paragraph(
            f"Confidence: {overall.get('confidence', 'Low')}"
        )

        document.add_paragraph(
            str(overall.get("reason", ""))
        )

    output = io.BytesIO()
    document.save(output)

    return output.getvalue()


EXPORTERS["DOCX"] = package_to_docx
