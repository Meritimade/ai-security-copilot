from __future__ import annotations

import base64
import io
import json
import os
import re
from pathlib import Path
from typing import Any

import pandas as pd
from dotenv import load_dotenv

load_dotenv()

STRUCTURED_EXTENSIONS = {
    ".csv",
    ".xlsx",
    ".xls",
    ".parquet",
    ".json",
}

TEXT_EXTENSIONS = {
    ".txt",
    ".log",
    ".rtf",
    ".json",
}

DOCUMENT_EXTENSIONS = {
    ".pdf",
    ".docx",
}

IMAGE_EXTENSIONS = {
    ".png",
    ".jpg",
    ".jpeg",
}

SUPPORTED_EXTENSIONS = (
    sorted(
        STRUCTURED_EXTENSIONS
        | TEXT_EXTENSIONS
        | DOCUMENT_EXTENSIONS
        | IMAGE_EXTENSIONS
    )
)


def get_openai_client():
    """Create an OpenAI client using Streamlit secrets or the local environment."""
    api_key = os.getenv("OPENAI_API_KEY")

    try:
        import streamlit as st

        try:
            secret_key = st.secrets.get("OPENAI_API_KEY")
            if secret_key:
                api_key = secret_key
        except Exception:
            pass
    except Exception:
        pass

    if not api_key:
        return None

    from openai import OpenAI

    return OpenAI(api_key=api_key)


def classify_file(filename: str) -> str:
    ext = Path(filename).suffix.lower()

    if ext in STRUCTURED_EXTENSIONS:
        return "structured"

    if ext in IMAGE_EXTENSIONS:
        return "image"

    if ext in DOCUMENT_EXTENSIONS or ext in TEXT_EXTENSIONS:
        return "document"

    return "unsupported"


def safe_filename(filename: str) -> str:
    """Prevent path traversal while retaining a readable filename."""
    name = Path(filename).name
    name = re.sub(r"[^A-Za-z0-9._ -]", "_", name)
    return name[:180] or "uploaded_file"


def read_structured_file(path: Path) -> pd.DataFrame:
    ext = path.suffix.lower()

    if ext == ".csv":
        return pd.read_csv(path, low_memory=False)

    if ext == ".xlsx":
        return pd.read_excel(path, engine="openpyxl")

    if ext == ".xls":
        return pd.read_excel(path, engine="xlrd")

    if ext == ".parquet":
        return pd.read_parquet(path)

    if ext == ".json":
        raw = json.loads(
            path.read_text(
                encoding="utf-8",
                errors="replace"
            )
        )

        if isinstance(raw, list):
            return pd.json_normalize(raw)

        if isinstance(raw, dict):
            # Prefer a record-list value when one exists.
            for value in raw.values():
                if isinstance(value, list):
                    return pd.json_normalize(value)

            return pd.json_normalize(raw)

        raise ValueError("The JSON structure could not be converted to a table.")

    raise ValueError(f"Unsupported structured format: {ext}")


def convert_structured_to_csv(source_path: Path, destination_path: Path) -> dict[str, Any]:
    df = read_structured_file(source_path)

    if df.empty:
        raise ValueError("The uploaded structured file contains no records.")

    destination_path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(destination_path, index=False)

    return {
        "source_format": source_path.suffix.lower().lstrip("."),
        "rows": int(len(df)),
        "columns": int(len(df.columns)),
        "column_names": [str(c) for c in df.columns],
        "processed_path": str(destination_path),
    }


def extract_pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""

        if text.strip():
            pages.append(f"[Page {page_number}]\n{text.strip()}")

    return "\n\n".join(pages)


def extract_docx_text(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    parts = []

    for paragraph in document.paragraphs:
        value = paragraph.text.strip()
        if value:
            parts.append(value)

    for table_index, table in enumerate(document.tables, start=1):
        parts.append(f"[Table {table_index}]")

        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))

    return "\n".join(parts)


def extract_rtf_text(path: Path) -> str:
    from striprtf.striprtf import rtf_to_text

    return rtf_to_text(
        path.read_text(
            encoding="utf-8",
            errors="replace"
        )
    )


def extract_text_file(path: Path) -> str:
    return path.read_text(
        encoding="utf-8",
        errors="replace"
    )


def extract_json_text(path: Path) -> str:
    raw = path.read_text(
        encoding="utf-8",
        errors="replace"
    )

    try:
        parsed = json.loads(raw)
        return json.dumps(parsed, indent=2, ensure_ascii=False)
    except Exception:
        return raw


def extract_document_text(path: Path) -> str:
    ext = path.suffix.lower()

    if ext == ".pdf":
        return extract_pdf_text(path)

    if ext == ".docx":
        return extract_docx_text(path)

    if ext == ".rtf":
        return extract_rtf_text(path)

    if ext == ".json":
        return extract_json_text(path)

    return extract_text_file(path)


def _clean_json_response(value: str) -> str:
    value = value.strip()

    if value.startswith("```"):
        value = re.sub(r"^```(?:json)?\s*", "", value)
        value = re.sub(r"\s*```$", "", value)

    return value.strip()


def _parse_ai_json(value: str) -> dict[str, Any]:
    cleaned = _clean_json_response(value)

    try:
        parsed = json.loads(cleaned)
        return parsed if isinstance(parsed, dict) else {}
    except Exception:
        match = re.search(r"\{.*\}", cleaned, flags=re.DOTALL)

        if not match:
            return {}

        try:
            parsed = json.loads(match.group(0))
            return parsed if isinstance(parsed, dict) else {}
        except Exception:
            return {}


def analyse_text_with_ai(
    text: str,
    filename: str,
    file_type: str,
    model: str = "gpt-5.6",
) -> dict[str, Any]:
    client = get_openai_client()

    if client is None:
        raise RuntimeError(
            "OPENAI_API_KEY is not configured. Add it to the local .env "
            "or Streamlit Cloud Secrets before analysing documents with AI."
        )

    # Keep very large documents manageable. The source file itself is preserved;
    # this limit only controls the amount of text sent to the model.
    max_chars = 120_000
    truncated = len(text) > max_chars
    analysis_text = text[:max_chars]

    prompt = f"""
You are the security analyst component of an AI Security Copilot.

Analyse the uploaded {file_type} security material below.

File: {filename}

Rules:
- Use only evidence present in the supplied material.
- Do not invent IPs, users, hosts, processes, commands, timestamps, attacks,
  vulnerabilities, or other facts.
- Distinguish observed evidence from inference.
- If the material does not establish an attack, say so.
- MITRE ATT&CK techniques are hypotheses unless the evidence supports them.
- Never treat a filename or dataset label alone as proof.
- Explicitly identify missing telemetry.
- Give practical SOC investigation and defensive actions.
- For narrative documents, do not pretend the document is network-flow telemetry.

Return valid JSON with this structure:

{{
  "executive_summary": "short factual summary",
  "document_type_assessment": "what the material appears to contain",
  "findings": [
    {{
      "finding": "name",
      "classification": "security_event|suspicious|informational|benign|unknown",
      "severity": "Critical|High|Medium|Low|Informational",
      "confidence": "High|Medium|Low",
      "observed_evidence": ["..."],
      "inference": ["..."],
      "evidence_gaps": ["..."],
      "investigation_steps": ["..."],
      "defensive_actions": ["..."],
      "mitre_candidates": [
        {{
          "technique_id": "Txxxx",
          "technique_name": "name",
          "assessment": "SUPPORTED|POSSIBLE|NOT_SUPPORTED",
          "reason": "brief evidence-based reason"
        }}
      ]
    }}
  ],
  "telemetry_capabilities": {{
    "source_identity": false,
    "destination_identity": false,
    "source_port": false,
    "destination_port": false,
    "username": false,
    "hostname": false,
    "process": false,
    "command": false,
    "authentication_outcome": false,
    "application_layer_events": false
  }},
  "overall_risk": {{
    "level": "Critical|High|Medium|Low|Informational|Unknown",
    "confidence": "High|Medium|Low",
    "reason": "brief evidence-based reason"
  }}
}}

Uploaded material:

--- BEGIN MATERIAL ---
{analysis_text}
--- END MATERIAL ---
"""

    response = client.responses.create(
        model=model,
        input=prompt,
    )

    output_text = getattr(response, "output_text", "") or ""
    result = _parse_ai_json(output_text)

    if not result:
        return {
            "executive_summary": output_text,
            "document_type_assessment": file_type,
            "findings": [],
            "telemetry_capabilities": {},
            "overall_risk": {
                "level": "Unknown",
                "confidence": "Low",
                "reason": "The AI response could not be parsed into the structured report format.",
            },
            "analysis_warning": "The model returned non-structured output.",
        }

    result["source_file"] = filename
    result["source_type"] = file_type
    result["text_truncated_for_ai"] = truncated
    result["characters_analysed"] = len(analysis_text)

    return result


def analyse_image_with_ai(
    path: Path,
    filename: str,
    model: str = "gpt-5.6",
) -> dict[str, Any]:
    client = get_openai_client()

    if client is None:
        raise RuntimeError(
            "OPENAI_API_KEY is not configured. Add it to the local .env "
            "or Streamlit Cloud Secrets before analysing images with AI."
        )

    mime = {
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
    }.get(path.suffix.lower(), "image/png")

    encoded = base64.b64encode(path.read_bytes()).decode("ascii")
    data_url = f"data:{mime};base64,{encoded}"

    prompt = f"""
Analyse this uploaded security-related image as an AI Security Copilot.

File: {filename}

The image may be a screenshot of a SIEM alert, dashboard, terminal,
security event, log, diagram, or other security material.

Rules:
- Read only what is visibly supported by the image.
- Do not invent values that cannot be read.
- Clearly distinguish observed evidence from inference.
- Do not identify a person from an image.
- Do not treat a screenshot label as proof of an attack.
- MITRE ATT&CK mappings are hypotheses unless the visible evidence supports them.
- Identify what additional telemetry an analyst should collect.

Return valid JSON using:
{{
  "executive_summary": "...",
  "document_type_assessment": "security screenshot|diagram|log|other",
  "findings": [
    {{
      "finding": "...",
      "classification": "security_event|suspicious|informational|benign|unknown",
      "severity": "Critical|High|Medium|Low|Informational",
      "confidence": "High|Medium|Low",
      "observed_evidence": ["..."],
      "inference": ["..."],
      "evidence_gaps": ["..."],
      "investigation_steps": ["..."],
      "defensive_actions": ["..."],
      "mitre_candidates": [
        {{
          "technique_id": "Txxxx",
          "technique_name": "...",
          "assessment": "SUPPORTED|POSSIBLE|NOT_SUPPORTED",
          "reason": "..."
        }}
      ]
    }}
  ],
  "overall_risk": {{
    "level": "Critical|High|Medium|Low|Informational|Unknown",
    "confidence": "High|Medium|Low",
    "reason": "..."
  }}
}}
"""

    response = client.responses.create(
        model=model,
        input=[
            {
                "role": "user",
                "content": [
                    {
                        "type": "input_text",
                        "text": prompt,
                    },
                    {
                        "type": "input_image",
                        "image_url": data_url,
                    },
                ],
            }
        ],
    )

    output_text = getattr(response, "output_text", "") or ""
    result = _parse_ai_json(output_text)

    if not result:
        result = {
            "executive_summary": output_text,
            "document_type_assessment": "image",
            "findings": [],
            "overall_risk": {
                "level": "Unknown",
                "confidence": "Low",
                "reason": "The AI response could not be parsed into the structured report format.",
            },
        }

    result["source_file"] = filename
    result["source_type"] = "image"

    return result


def analyse_uploaded_file(
    path: Path,
    filename: str | None = None,
) -> dict[str, Any]:
    filename = filename or path.name
    file_class = classify_file(filename)

    if file_class == "structured":
        return {
            "route": "structured",
            "filename": filename,
            "message": (
                "Structured security data should use the existing "
                "cleaning, validation, schema discovery, behavioural analysis, "
                "MITRE and risk pipeline."
            ),
        }

    if file_class == "document":
        extracted_text = extract_document_text(path)

        if not extracted_text.strip():
            return {
                "route": "document",
                "filename": filename,
                "source_type": Path(filename).suffix.lower(),
                "executive_summary": (
                    "No extractable text was found in this document. "
                    "The file may be scanned/image-only content and requires OCR "
                    "or vision processing."
                ),
                "findings": [],
                "overall_risk": {
                    "level": "Unknown",
                    "confidence": "Low",
                    "reason": "No machine-readable text was available.",
                },
                "evidence_gaps": [
                    "Machine-readable document text was unavailable."
                ],
            }

        result = analyse_text_with_ai(
            extracted_text,
            filename,
            Path(filename).suffix.lower().lstrip("."),
        )
        result["route"] = "document"
        return result

    if file_class == "image":
        result = analyse_image_with_ai(path, filename)
        result["route"] = "image"
        return result

    raise ValueError(
        f"Unsupported file type: {Path(filename).suffix.lower()}. "
        f"Supported types: {', '.join(SUPPORTED_EXTENSIONS)}"
    )
